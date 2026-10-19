"""Integration tests: ingest job (diff, upsert, merge, history, rollback) and row-level security.

Needs a Postgres server. Set TEST_DATABASE_URL_ADMIN to a superuser/owner URL whose database can be dropped
and recreated, e.g. postgresql://nazryx:pw@localhost:5432/nazryx_test. Skipped when not set.
"""
import importlib
import json
import os
import uuid

import psycopg
import pytest
from psycopg import sql

URL = os.environ.get("TEST_DATABASE_URL_ADMIN")
pytestmark = pytest.mark.skipif(not URL, reason="TEST_DATABASE_URL_ADMIN not set")

HEADER = ["No", "Product Category", "Certificate Number", "Brand Name", "Classification", "Generic Name", "Dosage Form",
          "National ID No", "Active Pharmaceutical Ingredients", "Product Strength", "Registrant", "Registrant Country",
          "Local Technical Represenatative", "Manufacturer", "Manufacturing Country", "Registration Status"]
MOLS = ["Amoxicillin", "Paracetamol", "Metformin Hydrochloride", "Ciprofloxacin", "Omeprazole", "Amlodipine",
        "Ceftriaxone Sodium", "Ibuprofen", "Sodium Chloride", "Fluconazole", "Atenolol", "Salbutamol"]
LTRS = ["ALPHA PHARMA LIMITED", "Alpha Pharma Ltd", "Beta Distributors (T) Ltd", "GAMMA HEALTHCARE LTD"]


def registry_rows(n=130, extra=(), cancel=()):
    rows = []
    for i in range(n):
        cert = f"TAN 2{i % 6} HM {i:04d}"
        status = "Cancelled/Withdrawn" if cert in cancel else "Registered/Compliant"
        rows.append([str(i + 1), "Medicines", cert, f"Brand{i}", "Human Medicinal Product", MOLS[i % len(MOLS)], "Tablet", "",
                     "", "500", f"Maker {i % 7} Ltd", "INDIA", LTRS[i % len(LTRS)], f"Maker {i % 7} Ltd", "INDIA", status])
    for j, (mol, ltr) in enumerate(extra):
        rows.append([str(n + j + 1), "Medicines", f"TAN 26 HM 9{j:03d}", f"New{j}", "Human Medicinal Product", mol, "Tablet",
                     "", "", "10", "Maker 1 Ltd", "INDIA", ltr, "Maker 1 Ltd", "INDIA", "Registered/Compliant"])
    return rows


def write_registry(path, rows):
    td = lambda v: f"<td>{v}</td>"
    html = ["<table border='1'><tr><td colspan='12'><b>Registered Products</b></td></tr>",
            "<tr>" + "".join(td(h) for h in HEADER) + "</tr>"]
    html += ["<tr>" + "".join(td(v) for v in r) + "</tr>" for r in rows]
    html.append("</table>")
    path.write_text("\n".join(html))
    return str(path)


NEMLIT = """The National Essential Medicines List 2026 Edition
6.0 Anti-Infective Medicines
Amoxicillin Capsule 250mg A
Ciprofloxacin Tablet 250mg B
Doxycycline Capsule 100mg A
Ceftriaxone Injection 1 g vial B
13.0 Cardiovascular Medicines
Amlodipine Tablet 5mg A
Atenolol Tablet 50mg A
Digoxin Tablet 250mcg c
"""
WHO = "amoxicillin                     Capsule\nciprofloxacin                   Tablet\ndoxycycline\ndigoxin\n"


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("ingest")
    base = psycopg.conninfo.conninfo_to_dict(URL)
    dbname = base.get("dbname") or "nazryx_test"
    with psycopg.connect(psycopg.conninfo.make_conninfo(URL, dbname="postgres"), autocommit=True) as c:
        c.execute(sql.SQL("drop database if exists {} with (force)").format(sql.Identifier(dbname)))
        c.execute(sql.SQL("create database {}").format(sql.Identifier(dbname)))
    (tmp / "nemlit.txt").write_text(NEMLIT)
    (tmp / "who.txt").write_text(WHO)
    os.environ.update(DATABASE_URL_ADMIN=URL, NEMLIT_FILE=str(tmp / "nemlit.txt"), WHO_FILE=str(tmp / "who.txt"),
                      APP_DB_PASSWORD="test-only-" + uuid.uuid4().hex[:8])
    import migrate
    import ingest_job
    importlib.reload(ingest_job)
    migrate.migrate()
    return {"tmp": tmp, "job": ingest_job}


def q(sql_text, *args):
    with psycopg.connect(URL) as c:
        return c.execute(sql_text, args).fetchall()


def test_initial_load_then_diffed_upload(env):
    tmp, job = env["tmp"], env["job"]
    first = write_registry(tmp / "r1.xls", registry_rows())
    job.run(job.create_job(first, "initial"))
    assert q("select count(*) from registrations where active")[0][0] == 130
    # the two spellings of Alpha Pharma are one company
    alpha = q("select count(distinct ltr_id) from registrations r join company_aliases a on a.company_id = r.ltr_id "
              "where a.alias in ('ALPHA PHARMA LIMITED', 'Alpha Pharma Ltd')")[0][0]
    assert alpha == 1
    # essential molecules with no registrations score highest; Digoxin is level C and unregistered
    top = q("select m.inn, g.score from gap_scores g join molecules m on m.id = g.molecule_id order by g.score desc, m.inn")
    assert top[0] == ("Doxycycline", 100) and ("Digoxin", 94) in [(n, float(s)) for n, s in top]
    before_scores = dict(q("select m.inn, g.score from gap_scores g join molecules m on m.id = g.molecule_id"))

    # second upload: 2 new doxycycline products (one from a new distributor), one product cancelled, one removed
    rows = registry_rows(n=129, extra=[("Doxycycline", "Delta Medics Ltd"), ("Doxycycline", "Beta Distributors (T) Ltd")],
                         cancel={"TAN 20 HM 0006"})
    second = write_registry(tmp / "r2.xls", rows)
    diff = job.run(job.create_job(second, "upload"))
    assert diff["new_registrations"] == 2
    assert diff["removed_or_cancelled"] == 2          # row 129 missing + TAN 20 HM 0006 cancelled
    assert diff["new_distributors"] == 1
    assert diff["gap_scores_moved"] >= 1
    after = dict(q("select m.inn, g.score from gap_scores g join molecules m on m.id = g.molecule_id"))
    assert float(after["Doxycycline"]) < float(before_scores["Doxycycline"])
    changes = dict(q("select change, count(*) from registration_history group by change"))
    assert changes["added"] == 2 and changes["removed"] == 1 and changes["status"] == 1
    assert q("select count(*) from feed_items where kind = 'registration'")[0][0] == 2


def test_failed_upload_rolls_back(env):
    tmp, job = env["tmp"], env["job"]
    before = q("select count(*), sum(id) from registrations where active")[0]
    gaps_before = q("select sum(score) from gap_scores")[0]
    # a partial export (far fewer rows than last time) must be rejected
    partial = write_registry(tmp / "partial.xls", registry_rows(n=110))
    rows_before = q("select count(*) from registrations")[0][0]
    job.MIN_RATIO = 0.9
    try:
        jid = job.create_job(partial, "upload")
        with pytest.raises(ValueError):
            job.run(jid)
    finally:
        job.MIN_RATIO = 0.6
    # a file that is not the registry at all
    bad = tmp / "bad.xls"
    bad.write_text("<table><tr><td>No</td><td>Something</td></tr><tr><td>1</td><td>x</td></tr></table>")
    jid2 = job.create_job(str(bad), "upload")
    with pytest.raises(Exception):
        job.run(jid2)
    status = dict(q("select id, status from registry_uploads where id in (%s, %s)", jid, jid2))
    assert status == {jid: "failed", jid2: "failed"}
    assert q("select count(*), sum(id) from registrations where active")[0] == before
    assert q("select count(*) from registrations")[0][0] == rows_before
    assert q("select sum(score) from gap_scores")[0] == gaps_before
    assert q("select error from registry_uploads where id = %s", jid)[0][0].startswith("Only")


def app_conn(env):
    info = psycopg.conninfo.conninfo_to_dict(URL)
    return psycopg.connect(psycopg.conninfo.make_conninfo(URL, user="nazryx_app", password=os.environ["APP_DB_PASSWORD"]))


def as_user(conn, uid):
    conn.execute("select set_config('app.user_id', %s, true)", (str(uid) if uid else "",))


def test_supplier_confidentiality_rls(env):
    with psycopg.connect(URL) as c:   # owner: create users, a supplier, an offer and two quotes
        cust = c.execute("insert into users (email, password_hash) values ('c@example.com', 'x') returning id").fetchone()[0]
        other = c.execute("insert into users (email, password_hash) values ('o@example.com', 'x') returning id").fetchone()[0]
        team = c.execute("insert into users (email, password_hash, role) values ('t@example.com', 'x', 'team') returning id").fetchone()[0]
        mol = c.execute("select id from molecules where inn = 'Doxycycline'").fetchone()[0]
        sup = c.execute("insert into suppliers (name, country) values ('Secret Supplier', 'IN') returning id").fetchone()[0]
        c.execute("insert into supplier_notes (supplier_id, notes) values (%s, 'internal')", (sup,))
        sup2 = c.execute("insert into suppliers (name, country) values ('Other Supplier', 'CN') returning id").fetchone()[0]
        off = c.execute("insert into supplier_offers (supplier_id, molecule_id, price) values (%s, %s, 1.25) returning id", (sup, mol)).fetchone()[0]
        off2 = c.execute("insert into supplier_offers (supplier_id, molecule_id, price) values (%s, %s, 0.99) returning id", (sup2, mol)).fetchone()[0]
        req = c.execute("insert into requests (user_id, molecule_id, quantity) values (%s, %s, '10000') returning id", (cust, mol)).fetchone()[0]
        q1 = c.execute("insert into quotes (request_id, supplier_offer_id, label) values (%s, %s, 'Supplier A') returning id", (req, off)).fetchone()[0]
        c.execute("insert into quotes (request_id, supplier_offer_id, label) values (%s, %s, 'Supplier B')", (req, off2))
        c.commit()

    with app_conn(env) as a:
        as_user(a, cust)
        assert a.execute("select count(*) from suppliers").fetchone()[0] == 0
        assert a.execute("select count(*) from supplier_offers").fetchone()[0] == 0
        assert a.execute("select label from quotes order by label").fetchall() == [("Supplier A",), ("Supplier B",)]
        assert a.execute("select count(*) from accepted_supplier").fetchone()[0] == 0
        assert a.execute("select molecule_supply_confirmed(%s)", (mol,)).fetchone()[0] is None   # free plan: hidden
        # RLS: the team-only update policy matches no rows for a customer, so nothing changes
        assert a.execute("update quotes set accepted = true where id = %s", (q1,)).rowcount == 0
        assert a.execute("update supplier_offers set price = 0 where id = %s", (off,)).rowcount == 0
        a.rollback()

        as_user(a, other)   # another customer cannot accept someone else's quote
        with pytest.raises(psycopg.errors.RaiseException):
            a.execute("select accept_quote(%s)", (q1,))
        a.rollback()

        as_user(a, cust)
        a.execute("select accept_quote(%s)", (q1,))
        assert a.execute("select name from suppliers").fetchall() == [("Secret Supplier",)]   # only the accepted one
        assert a.execute("select price from supplier_offers").fetchall()[0][0] == 1.25
        assert a.execute("select supplier_name, price from accepted_supplier").fetchall() == [("Secret Supplier", 1.25)]
        assert a.execute("select count(*) from supplier_notes").fetchone()[0] == 0
        assert a.execute("select stage from requests where id = %s", (req,)).fetchone()[0] == "Registration"   # accepted: on to registration
        a.commit()

        as_user(a, other)
        assert a.execute("select count(*) from suppliers").fetchone()[0] == 0
        assert a.execute("select count(*) from quotes").fetchone()[0] == 0
        assert a.execute("select count(*) from requests").fetchone()[0] == 0
        a.rollback()

        as_user(a, team)
        assert a.execute("select count(*) from suppliers").fetchone()[0] == 2
        assert a.execute("select count(*) from supplier_notes").fetchone()[0] == 1
        a.rollback()

        as_user(a, None)   # signed out sees nothing private
        assert a.execute("select count(*) from quotes").fetchone()[0] == 0
        assert a.execute("select count(*) from users").fetchone()[0] == 0


def test_free_plan_limits_in_database(env):
    with psycopg.connect(URL) as c:
        u = c.execute("insert into users (email, password_hash, credits) values ('f@example.com', 'x', 1) returning id").fetchone()[0]
        mols = [r[0] for r in c.execute("select id from molecules order by id limit 6").fetchall()]
        comps = [r[0] for r in c.execute("select id from companies order by id limit 4").fetchall()]
        c.commit()
    with app_conn(env) as a:
        as_user(a, u)
        assert a.execute("select enrich_molecule(%s)", (mols[0],)).fetchone()[0] == 0
        assert a.execute("select enrich_molecule(%s)", (mols[0],)).fetchone()[0] == 0     # re-open is free
        with pytest.raises(psycopg.errors.RaiseException):
            a.execute("select enrich_molecule(%s)", (mols[1],))
        a.rollback()
        as_user(a, u)
        for i in range(5):
            a.execute("insert into requests (user_id, molecule_id) values (%s, %s)", (u, mols[i]))
        with pytest.raises(psycopg.errors.RaiseException):
            a.execute("insert into requests (user_id, molecule_id) values (%s, %s)", (u, mols[5]))
        a.rollback()
        as_user(a, u)
        for cid in comps[:3]:
            a.execute("insert into follows (user_id, company_id) values (%s, %s)", (u, cid))
        with pytest.raises(psycopg.errors.RaiseException):
            a.execute("insert into follows (user_id, company_id) values (%s, %s)", (u, comps[3]))
        a.rollback()
        as_user(a, u)   # customers cannot change their own plan or credits
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            a.execute("update users set plan = 'paid' where id = %s", (u,))


def test_48_hour_reminder(env):
    import reminder
    with psycopg.connect(URL) as c:
        cur = c.cursor()
        cur.execute("update registry_uploads set finished_at = now() - interval '20 hours' where status = 'applied'")
        assert reminder.status(cur)["overdue"] is False
        assert reminder.check_and_remind(cur) is False
        cur.execute("update registry_uploads set finished_at = now() - interval '50 hours' where status = 'applied'")
        st = reminder.status(cur)
        assert st["overdue"] and 49 < st["hours_since"] < 51
        before = cur.execute("select count(*) from notifications where user_id is null").fetchone()[0]
        assert reminder.check_and_remind(cur) is True             # email + WhatsApp to the team (stubbed)
        chans = cur.execute("select channel from notifications where user_id is null order by id desc limit 2").fetchall()
        assert {c[0] for c in chans} == {"email", "whatsapp"}
        assert cur.execute("select count(*) from notifications where user_id is null").fetchone()[0] == before + 2
        assert reminder.check_and_remind(cur) is False            # not again within the same 48-hour window
        c.rollback()


def test_signin_throttle_and_team_password_reset(env):
    with app_conn(env) as a:
        as_user(a, None)
        for _ in range(4):
            a.execute("select auth_record_failure('slow@example.com')")
        assert a.execute("select auth_throttled('Slow@Example.com')").fetchone()[0] is False
        a.execute("select auth_record_failure('slow@example.com')")
        assert a.execute("select auth_throttled('slow@example.com')").fetchone()[0] is True
        assert a.execute("select auth_throttled('other@example.com')").fetchone()[0] is False
        a.execute("select auth_clear_failures('slow@example.com')")
        assert a.execute("select auth_throttled('slow@example.com')").fetchone()[0] is False
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            a.execute("select * from login_failures")
        a.rollback()
    with psycopg.connect(URL) as c:
        cust = c.execute("insert into users (email, password_hash) values ('reset@example.com', 'old') returning id").fetchone()[0]
        team = c.execute("select id from users where role = 'team' limit 1").fetchone()[0]
        c.execute("insert into sessions (token_hash, user_id, expires_at) values ('h1', %s, now() + interval '1 day')", (cust,))
        c.commit()
    with app_conn(env) as a:
        as_user(a, cust)   # a customer cannot reset anyone's password
        with pytest.raises(psycopg.errors.RaiseException):
            a.execute("select admin_reset_password(%s, 'x')", (cust,))
        a.rollback()
        as_user(a, team)
        a.execute("select admin_reset_password(%s, 'newhash')", (cust,))
        a.commit()
    assert q("select password_hash from users where id = %s", cust)[0][0] == "newhash"
    assert q("select count(*) from sessions where user_id = %s", cust)[0][0] == 0


def test_confirmed_supply_does_not_change_the_score(env):
    import loader
    with psycopg.connect(URL) as c:   # the RLS test added a confirmed Doxycycline offer
        mol = c.execute("select id from molecules where inn = 'Doxycycline'").fetchone()[0]
        before = c.execute("select score from gap_scores where molecule_id = %s", (mol,)).fetchone()[0]
        loader.recompute_gaps(c.cursor(), "TZ")
        after = c.execute("select score from gap_scores where molecule_id = %s", (mol,)).fetchone()[0]
        free = c.execute("insert into users (email, password_hash) values ('free2@example.com', 'x') returning id").fetchone()[0]
        paid = c.execute("insert into users (email, password_hash, plan) values ('paid2@example.com', 'x', 'paid') returning id").fetchone()[0]
        c.commit()
    assert before == after
    with app_conn(env) as a:
        seen = []
        for uid in (free, paid):
            as_user(a, uid)
            seen.append(a.execute("select score from visible_gap_scores where molecule_id = %s", (mol,)).fetchone()[0])
        assert seen == [after, after]
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            a.execute("select score from gap_scores")
        a.rollback()


def test_first_upload_into_empty_registry_is_the_baseline(env):
    """A first load sent as a normal upload must not announce every product in the market as new."""
    tmp, job = env["tmp"], env["job"]
    url = psycopg.conninfo.make_conninfo(URL, dbname="nazryx_test_first")
    with psycopg.connect(psycopg.conninfo.make_conninfo(URL, dbname="postgres"), autocommit=True) as c:
        c.execute("drop database if exists nazryx_test_first with (force)")
        c.execute("create database nazryx_test_first")
    import migrate
    os.environ["DATABASE_URL_ADMIN"] = url
    try:
        migrate.migrate()
        diff = job.run(job.create_job(write_registry(tmp / "first.xls", registry_rows()), "upload"))
        with psycopg.connect(url) as c:
            assert c.execute("select count(*) from registrations where active").fetchone()[0] == 130
            assert c.execute("select count(*) from feed_items where kind = 'registration'").fetchone()[0] == 0
            assert c.execute("select count(*) from registration_history").fetchone()[0] == 0
        assert diff["new_distributors"] == 0
    finally:
        os.environ["DATABASE_URL_ADMIN"] = URL
        with psycopg.connect(psycopg.conninfo.make_conninfo(URL, dbname="postgres"), autocommit=True) as c:
            c.execute("drop database if exists nazryx_test_first with (force)")


def test_stored_files_follow_request_access(env):
    with psycopg.connect(URL) as c:
        owner = c.execute("insert into users (email, password_hash) values ('files1@example.com', 'x') returning id").fetchone()[0]
        other = c.execute("insert into users (email, password_hash) values ('files2@example.com', 'x') returning id").fetchone()[0]
        req = c.execute("insert into requests (user_id, quantity) values (%s, '1') returning id", (owner,)).fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, owner)
        fid = a.execute("insert into stored_files (name, size, data, uploaded_by) values ('spec.pdf', 3, 'abc', %s) returning id",
                        (owner,)).fetchone()[0]
        a.execute("insert into request_files (request_id, path, file_name, uploaded_by) values (%s, %s, 'spec.pdf', %s)",
                  (req, f"db:{fid}", owner))
        a.commit()
        as_user(a, owner)
        assert bytes(a.execute("select data from stored_files where id = %s", (fid,)).fetchone()[0]) == b"abc"
        a.rollback()
        as_user(a, other)   # another customer cannot read it, even knowing the id
        assert a.execute("select count(*) from stored_files where id = %s", (fid,)).fetchone()[0] == 0
        with pytest.raises(psycopg.errors.InsufficientPrivilege):   # nor upload in someone else's name
            a.execute("insert into stored_files (name, size, data, uploaded_by) values ('x', 1, 'x', %s)", (owner,))
        a.rollback()


def test_reference_lists_from_database(env):
    import storage
    tmp = env["tmp"]
    (tmp / "nemlit2.txt").write_text(NEMLIT)
    with psycopg.connect(URL) as c:
        cur = c.cursor()
        storage.set_reference(cur, "national_list", str(tmp / "nemlit2.txt"))
        local = storage.reference_file(cur, "national_list", None, str(tmp))
        assert open(local).read() == NEMLIT
        assert storage.reference_file(cur, "global_list", None) is None   # not uploaded, no fallback
        c.rollback()


def test_team_actions_are_logged_customer_actions_are_not(env):
    with psycopg.connect(URL) as c:
        cust = c.execute("insert into users (email, password_hash) values ('audit1@example.com', 'x') returning id").fetchone()[0]
        team = c.execute("insert into users (email, password_hash, role) values ('audit-team@example.com', 'x', 'team') returning id").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, cust)
        rid = a.execute("insert into requests (user_id, quantity) values (%s, '5') returning id", (cust,)).fetchone()[0]
        a.commit()
        as_user(a, team)
        a.execute("update requests set stage = 'Regulatory check' where id = %s", (rid,))
        a.execute("select admin_reset_password(%s, 'h2')", (cust,))
        a.commit()
        as_user(a, team)
        rows = a.execute("select table_name, changes from audit_log where row_id in (%s, %s) order by id", (str(rid), str(cust))).fetchall()
        assert ("requests", {"stage": ["Product request", "Regulatory check"]}) in rows
        assert any(t == "users" and ch.get("password") == "reset" and "password_hash" not in ch for t, ch in rows)
        assert not any(t == "requests" and "quantity" in ch for t, ch in rows)    # the customer's insert is not logged
        a.rollback()
        as_user(a, cust)   # customers cannot read the log
        assert a.execute("select count(*) from audit_log").fetchone()[0] == 0


def test_signup_rate_and_upload_limits(env):
    with app_conn(env) as a:
        as_user(a, None)
        allowed = [a.execute("select rate_allow('signup', '10.0.0.9', 5, interval '1 hour')").fetchone()[0] for _ in range(6)]
        assert allowed == [True] * 5 + [False]
        assert a.execute("select rate_allow('signup', '10.0.0.10', 5, interval '1 hour')").fetchone()[0] is True
        a.commit()
    with psycopg.connect(URL) as c:
        u = c.execute("insert into users (email, password_hash) values ('uploads@example.com', 'x') returning id").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, u)
        for i in range(60):
            a.execute("insert into stored_files (name, size, data, uploaded_by) values ('f', 1, 'x', %s)", (u,))
        with pytest.raises(psycopg.errors.RaiseException, match="60 files a day"):
            a.execute("insert into stored_files (name, size, data, uploaded_by) values ('f', 1, 'x', %s)", (u,))
        a.rollback()


def test_two_step_challenge(env):
    with psycopg.connect(URL) as c:
        u = c.execute("insert into users (email, password_hash) values ('2step@example.com', 'x') returning id").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, u)
        a.execute("insert into user_totp (user_id, secret) values (%s, 'JBSWY3DPEHPK3PXP')", (u,))
        a.commit()
        as_user(a, None)
        assert a.execute("select auth_totp_enabled(%s)", (u,)).fetchone()[0] is False      # not confirmed yet
        a.commit()
        as_user(a, u)
        a.execute("update user_totp set enabled = true where user_id = %s", (u,))
        a.commit()
        as_user(a, None)
        assert a.execute("select auth_totp_enabled(%s)", (u,)).fetchone()[0] is True
        a.execute("select auth_create_challenge('tok', %s)", (u,))
        assert a.execute("select secret from auth_challenge('tok')").fetchone()[0] == "JBSWY3DPEHPK3PXP"
        for _ in range(5):
            a.execute("select auth_challenge_failed('tok')")
        assert a.execute("select count(*) from auth_challenge('tok')").fetchone()[0] == 0      # 5 wrong codes: dead
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            a.execute("select * from login_challenges")
        a.rollback()
        as_user(a, None)   # nobody else reads someone's secret
        assert a.execute("select count(*) from user_totp").fetchone()[0] == 0


def test_evaluations_hidden_picks_and_moves(env):
    with psycopg.connect(URL) as c:
        cust = c.execute("insert into users (email, password_hash) values ('ev1@example.com', 'x') returning id").fetchone()[0]
        other = c.execute("insert into users (email, password_hash) values ('ev2@example.com', 'x') returning id").fetchone()[0]
        team = c.execute("insert into users (email, password_hash, role) values ('evt@example.com', 'x', 'team') returning id").fetchone()[0]
        mol = c.execute("select id from molecules where inn = 'Doxycycline'").fetchone()[0]
        req = c.execute("insert into requests (user_id, molecule_id) values (%s, %s) returning id", (cust, mol)).fetchone()[0]
        c.execute("insert into portfolio_items (user_id, molecule_id) values (%s, %s)", (cust, mol))
        c.commit()
    with app_conn(env) as a:
        as_user(a, cust)   # a customer cannot write an evaluation, even on their own request
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            a.execute("insert into request_evaluations (request_id, verdict, summary) values (%s, 'go', 'x')", (req,))
        a.rollback()
        as_user(a, team)
        a.execute("insert into request_evaluations (request_id, verdict, summary, price_range) values (%s, 'go', 'Worth it', 'USD 1')", (req,))
        a.commit()
        as_user(a, cust)
        assert a.execute("select verdict, price_range from request_evaluations").fetchall() == [("go", "USD 1")]
        a.execute("insert into hidden_molecules (user_id, molecule_id) values (%s, %s)", (cust, mol))
        # the earlier diffed upload added Doxycycline products: they show as moves on this customer's portfolio
        moves = a.execute("select change, molecule from my_market_moves(3650)").fetchall()
        assert moves and all(m == "Doxycycline" for _, m in moves)
        a.commit()
        as_user(a, other)
        assert a.execute("select count(*) from request_evaluations").fetchone()[0] == 0
        assert a.execute("select count(*) from hidden_molecules").fetchone()[0] == 0
        assert a.execute("select count(*) from my_market_moves(3650)").fetchone()[0] == 0   # no portfolio, no moves
        a.rollback()
    assert q("select count(*) from audit_log where table_name = 'request_evaluations'")[0][0] == 1


def test_sourcing_list_and_free_cap_counts_only_early_stages(env):
    with psycopg.connect(URL) as c:
        cust = c.execute("insert into users (email, password_hash) values ('list1@example.com', 'x') returning id").fetchone()[0]
        other = c.execute("insert into users (email, password_hash) values ('list2@example.com', 'x') returning id").fetchone()[0]
        mol = c.execute("select id from molecules where inn = 'Doxycycline'").fetchone()[0]
        for stage in ("Registration", "Shipment", "Closed"):   # products past pricing do not count toward the free cap
            c.execute("insert into requests (user_id, molecule_id, stage) values (%s, %s, %s)", (cust, mol, stage))
        c.commit()
    with app_conn(env) as a:
        as_user(a, cust)
        a.execute("insert into sourcing_list (user_id, molecule_id) values (%s, %s)", (cust, mol))
        with pytest.raises(psycopg.errors.InsufficientPrivilege):   # a customer's list is their own
            a.execute("insert into sourcing_list (user_id, molecule_id) values (%s, %s)", (other, mol))
        a.rollback()
        as_user(a, cust)
        for _ in range(5):
            a.execute("insert into requests (user_id, molecule_id) values (%s, %s)", (cust, mol))
        with pytest.raises(psycopg.errors.RaiseException):          # the sixth early-stage request is over the free cap
            a.execute("insert into requests (user_id, molecule_id) values (%s, %s)", (cust, mol))
        a.rollback()
        as_user(a, other)
        assert a.execute("select count(*) from sourcing_list").fetchone()[0] == 0


def test_dossier_and_shipment_documents_are_team_written(env):
    with psycopg.connect(URL) as c:
        cust = c.execute("insert into users (email, password_hash) values ('dos1@example.com', 'x') returning id").fetchone()[0]
        other = c.execute("insert into users (email, password_hash) values ('dos2@example.com', 'x') returning id").fetchone()[0]
        team = c.execute("insert into users (email, password_hash, role) values ('dost@example.com', 'x', 'team') returning id").fetchone()[0]
        mol = c.execute("select id from molecules where inn = 'Doxycycline'").fetchone()[0]
        req = c.execute("insert into requests (user_id, molecule_id) values (%s, %s) returning id", (cust, mol)).fetchone()[0]
        sup = c.execute("insert into suppliers (name) values ('Dossier Supplier') returning id").fetchone()[0]
        off = c.execute("insert into supplier_offers (supplier_id, molecule_id) values (%s, %s) returning id", (sup, mol)).fetchone()[0]
        q = c.execute("""insert into quotes (request_id, supplier_offer_id, label, payment_terms, shelf_life_pct, gmp_by, supplier_type)
                         values (%s, %s, 'Supplier A', '30%% deposit', 80, 'WHO prequalified', 'manufacturer') returning id""", (req, off)).fetchone()[0]
        order = c.execute("insert into orders (request_id, quote_id, status) values (%s, %s, 'Import permit') returning id", (req, q)).fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, team)
        a.execute("insert into dossier_items (request_id, item, status) values (%s, 'gmp', 'have')", (req,))
        a.execute("insert into request_dossier (request_id, be_status) values (%s, 'needed')", (req,))
        a.execute("insert into shipment_documents (order_id, doc, ready) values (%s, 'permit', true)", (order,))
        a.commit()
        as_user(a, cust)
        assert a.execute("select status from dossier_items").fetchall() == [("have",)]
        assert a.execute("select be_status from request_dossier").fetchall() == [("needed",)]
        assert a.execute("select doc from shipment_documents").fetchall() == [("permit",)]
        assert a.execute("select shelf_life_pct, supplier_type from quotes where id = %s", (q,)).fetchone() == (80, "manufacturer")
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            a.execute("insert into dossier_items (request_id, item, status) values (%s, 'cpp', 'have')", (req,))
        a.rollback()
        as_user(a, other)
        assert a.execute("select count(*) from dossier_items").fetchone()[0] == 0
        assert a.execute("select count(*) from shipment_documents").fetchone()[0] == 0
        a.rollback()


def test_focus_questions_and_routes(env):
    with psycopg.connect(URL) as c:
        free = c.execute("insert into users (email, password_hash) values ('fq1@example.com', 'x') returning id").fetchone()[0]
        other = c.execute("insert into users (email, password_hash) values ('fq2@example.com', 'x') returning id").fetchone()[0]
        paid = c.execute("insert into users (email, password_hash, plan) values ('fq3@example.com', 'x', 'paid') returning id").fetchone()[0]
        mol = c.execute("select id from molecules where inn = 'Doxycycline'").fetchone()[0]
        c.execute("insert into molecule_routes (molecule_id, route_key, note) values (%s, 'abridged', 'WHO PQ reference exists')", (mol,))
        c.commit()
    with app_conn(env) as a:
        as_user(a, free)
        a.execute("insert into user_focus (user_id, classes, goals) values (%s, %s, %s)", (free, ["Anti-infective Medicines"], ["new_molecules"]))
        for i in range(3):
            a.execute("insert into expert_questions (user_id, body) values (%s, %s)", (free, f"q{i}"))
        with pytest.raises(psycopg.errors.RaiseException):          # the fourth question this month is over the free limit
            a.execute("insert into expert_questions (user_id, body) values (%s, 'q4')", (free,))
        a.rollback()
        as_user(a, free)
        with pytest.raises(psycopg.errors.InsufficientPrivilege):   # customers cannot answer their own questions
            a.execute("insert into expert_questions (user_id, body, answer) values (%s, 'x', 'self answer')", (free,))
        a.rollback()
        as_user(a, free)
        assert a.execute("select count(*) from molecule_routes").fetchone()[0] == 0     # route detail is paid
        assert a.execute("select count(*) from registration_routes").fetchone()[0] == 3
        as_user(a, other)
        assert a.execute("select count(*) from user_focus").fetchone()[0] == 0
        assert a.execute("select count(*) from expert_questions").fetchone()[0] == 0
        as_user(a, paid)
        assert a.execute("select route_key from molecule_routes where molecule_id = %s", (mol,)).fetchone()[0] == "abridged"
        a.rollback()


def test_manufacturer_side_privacy(env):
    with psycopg.connect(URL) as c:
        maker = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('mk1@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        other = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('mk2@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        team = c.execute("insert into users (email, password_hash, role) values ('mkt@example.com', 'x', 'team') returning id").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, maker)
        a.execute("insert into manufacturer_profiles (user_id, company_name) values (%s, 'Maker One')", (maker,))
        pid = a.execute("insert into mfr_products (user_id, name) values (%s, 'Doxycycline 100mg capsules') returning id", (maker,)).fetchone()[0]
        eid = a.execute("insert into engagements (user_id, kind, country, product_id) values (%s, 'market_entry', 'TZ', %s) returning id", (maker, pid)).fetchone()[0]
        a.commit()
        as_user(a, maker)
        with pytest.raises(psycopg.errors.RaiseException):          # cannot approve itself
            a.execute("update manufacturer_profiles set status = 'approved' where user_id = %s", (maker,))
        a.rollback()
        as_user(a, maker)
        with pytest.raises(psycopg.errors.RaiseException):          # cannot switch account type after onboarding
            a.execute("update users set account_type = 'distributor' where id = %s", (maker,))
        a.rollback()
        as_user(a, maker)
        assert a.execute("update engagements set stage = 'Approved' where id = %s", (eid,)).rowcount == 0   # only the team moves stages
        profiles = a.execute("select * from partner_profiles('Anti-infective Medicines', null)").fetchall()
        assert profiles and all(len(p[0]) == 8 for p in profiles)                                        # codes, no names
        with pytest.raises(psycopg.errors.RaiseException):          # only the team turns a code into a company
            a.execute("select partner_company(%s)", (profiles[0][0],))
        a.rollback()
        as_user(a, other)
        assert a.execute("select count(*) from mfr_products").fetchone()[0] == 0
        assert a.execute("select count(*) from engagements").fetchone()[0] == 0
        assert a.execute("select count(*) from manufacturer_profiles").fetchone()[0] == 0
        a.rollback()
        as_user(a, team)
        code = a.execute("select code from partner_profiles('Anti-infective Medicines', null) limit 1").fetchone()[0]
        assert a.execute("select partner_company(%s)", (code,)).fetchone()[0] is not None
        a.execute("update manufacturer_profiles set status = 'approved' where user_id = %s", (maker,))
        a.execute("update engagements set stage = 'Market assessment' where id = %s", (eid,))
        a.commit()
    assert q("select status from manufacturer_profiles where user_id = %s", maker)[0][0] == "approved"


def test_team_documents_everywhere_and_internal_files(env):
    with psycopg.connect(URL) as c:
        cust = c.execute("insert into users (email, password_hash, onboarded) values ('doc1@example.com', 'x', true) returning id").fetchone()[0]
        other = c.execute("insert into users (email, password_hash, onboarded) values ('doc2@example.com', 'x', true) returning id").fetchone()[0]
        team = c.execute("insert into users (email, password_hash, role) values ('doct@example.com', 'x', 'team') returning id").fetchone()[0]
        rid = c.execute("insert into requests (user_id, molecule_text) values (%s, 'Doc test') returning id", (cust,)).fetchone()[0]
        sid = c.execute("insert into suppliers (name) values ('Doc Supplier') returning id").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, team)
        f1 = a.execute("insert into stored_files (name, size, data, uploaded_by) values ('a.pdf', 1, 'x', %s) returning id", (team,)).fetchone()[0]
        f2 = a.execute("insert into stored_files (name, size, data, uploaded_by) values ('b.pdf', 1, 'x', %s) returning id", (team,)).fetchone()[0]
        f3 = a.execute("insert into stored_files (name, size, data, uploaded_by) values ('c.pdf', 1, 'x', %s) returning id", (team,)).fetchone()[0]
        a.execute("""insert into request_files (request_id, path, file_name, kind, item, from_team, internal, uploaded_by)
                     values (%s, %s, 'gmp.pdf', 'dossier', 'gmp', true, false, %s), (%s, %s, 'margin.pdf', 'other', null, true, true, %s)""",
                  (rid, f"db:{f1}", team, rid, f"db:{f2}", team))
        a.execute("""insert into account_documents (user_id, kind, file_path, file_name, internal, uploaded_by)
                     values (%s, 'agreement', %s, 'agreement.pdf', false, %s), (%s, 'other', %s, 'notes.pdf', true, %s)""",
                  (cust, f"db:{f3}", team, cust, f"db:{f2}", team))
        a.execute("insert into supplier_documents (supplier_id, kind, file_path, file_name, uploaded_by) values (%s, 'gmp', %s, 'gmp.pdf', %s)",
                  (sid, f"db:{f1}", team))
        a.execute("insert into plan_prices (key, price) values ('mfr_entry', 'USD 1 / month') on conflict (key) do update set price = excluded.price")
        a.commit()
        as_user(a, cust)
        assert [r[0] for r in a.execute("select file_name from request_files where request_id = %s", (rid,))] == ["gmp.pdf"]   # internal hidden
        assert [r[0] for r in a.execute("select file_name from account_documents")] == ["agreement.pdf"]
        assert a.execute("select count(*) from supplier_documents").fetchone()[0] == 0                                   # team only
        assert a.execute("select count(*) from stored_files where id = %s", (f2,)).fetchone()[0] == 0                    # internal file body too
        assert a.execute("select count(*) from stored_files where id = %s", (f3,)).fetchone()[0] == 1
        assert a.execute("select price from plan_prices where key = 'mfr_entry'").fetchone()[0] == "USD 1 / month"
        assert a.execute("delete from request_files where request_id = %s", (rid,)).rowcount == 0                       # only the team removes
        with pytest.raises(psycopg.errors.InsufficientPrivilege):                                                       # cannot pose as the team
            a.execute("""insert into request_files (request_id, path, file_name, kind, from_team, uploaded_by)
                         values (%s, %s, 'fake.pdf', 'dossier', true, %s)""", (rid, f"db:{f3}", cust))
        a.rollback()
        as_user(a, other)
        assert a.execute("select count(*) from account_documents").fetchone()[0] == 0
        assert a.execute("select count(*) from request_files where request_id = %s", (rid,)).fetchone()[0] == 0
        a.rollback()
        as_user(a, team)
        assert a.execute("delete from request_files where request_id = %s and internal", (rid,)).rowcount == 1
        a.commit()
    assert q("select count(*) from audit_log where table_name = 'request_files' and action = 'delete'")[0][0] >= 1


def test_manufacturer_onboarding_answers(env):
    with psycopg.connect(URL) as c:
        maker = c.execute("insert into users (email, password_hash, account_type) values ('mk9@example.com', 'x', 'manufacturer') returning id").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, maker)
        a.execute("insert into manufacturer_profiles (user_id, company_name) values (%s, 'Maker Nine')", (maker,))
        a.execute("""update manufacturer_profiles set dosage_forms = '{Tablets}', target_countries = '{TZ,KE,UZ}', timeline = 'now',
                     goals = '{launch}', services = '{register,sales}', launch_scope = 'data_first', step = 6, suggested_plan = 'entry', chosen_plan = 'entry',
                     completed_at = now() where user_id = %s""", (maker,))
        a.execute("update users set onboarded = true where id = %s", (maker,))
        a.commit()
        as_user(a, maker)
        with pytest.raises(psycopg.errors.RaiseException):          # still cannot approve itself
            a.execute("update manufacturer_profiles set status = 'approved' where user_id = %s", (maker,))
        a.rollback()
    assert q("select chosen_plan, step from manufacturer_profiles where user_id = %s", maker)[0] == ("entry", 6)


def test_plan_tier_partner_search_and_catalog(env):
    with psycopg.connect(URL) as c:
        maker = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('mk7@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        team = c.execute("insert into users (email, password_hash, role) values ('mk7t@example.com', 'x', 'team') returning id").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, maker)
        a.execute("insert into manufacturer_profiles (user_id, company_name) values (%s, 'Maker Seven')", (maker,))
        mol = a.execute("""select molecule_id from registrations where active and country = 'TZ' group by 1 order by count(*) desc limit 1""").fetchone()[0]
        a.execute("insert into mfr_products (user_id, name, molecule_id) values (%s, 'Top molecule', %s)", (maker, mol))
        a.commit()
        as_user(a, maker)
        assert a.execute("select mfr_tier()").fetchone()[0] == 0
        rows = a.execute("select code, open_to, carries_mine, total, mine from partner_search('TZ')").fetchall()
        assert 0 < len(rows) <= 3 and rows[0][3] >= len(rows)                     # no plan: three rows, total still known
        assert all(r[1] + r[2] == r[4] == 1 for r in rows)                          # open to + carries = my molecules
        with pytest.raises(psycopg.errors.RaiseException):                          # cannot give itself a plan
            a.execute("update manufacturer_profiles set active_plan = 'entry' where user_id = %s", (maker,))
        a.rollback()
        assert a.execute("select count(*) from medicine_catalog").fetchone()[0] >= 0  # readable by everyone signed in
        a.rollback()
        as_user(a, team)
        a.execute("update manufacturer_profiles set active_plan = 'data', active_since = now() where user_id = %s", (maker,))
        a.commit()
        as_user(a, maker)
        assert a.execute("select mfr_tier()").fetchone()[0] == 1
        target = a.execute("select total from partner_search('TZ')").fetchall()
        assert len(target) == min(target[0][0], 10)                                  # Basic Market Data: a target list of ten
        assert q("select plan from users where id = %s", maker)[0][0] == "paid"     # an active plan makes the account paid
        a.rollback()
        as_user(a, team)
        a.execute("update manufacturer_profiles set active_plan = 'access' where user_id = %s", (maker,))
        a.commit()
        as_user(a, maker)
        full = a.execute("select total from partner_search('TZ')").fetchall()
        assert len(full) == min(full[0][0], 200)                                     # Complete Market Access: every distributor
        a.rollback()


def test_payments_invoices_launches_and_watch(env):
    import billing
    with psycopg.connect(URL) as c:
        maker = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('pay1@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        other = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('pay2@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        team = c.execute("insert into users (email, password_hash, role) values ('payt@example.com', 'x', 'team') returning id").fetchone()[0]
        c.execute("insert into manufacturer_profiles (user_id, company_name) values (%s, 'Pay Pharma'), (%s, 'Other Pharma')", (maker, other))
        c.commit()
    with app_conn(env) as a:
        # the webhook (no signed-in user) records a payment; a signed-in customer cannot
        a.execute("""insert into payment_events (stripe_event_id, kind, user_id, plan, amount_cents, stripe_customer, stripe_subscription)
                     values ('evt_1', 'checkout_completed', %s, 'data', 19900, 'cus_1', 'sub_1'),
                            ('evt_2', 'invoice_paid', %s, 'data', 19900, 'cus_1', 'sub_1')""", (maker, maker))
        a.commit()
        as_user(a, maker)
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            a.execute("insert into payment_events (kind, user_id, plan) values ('checkout_completed', %s, 'entry')", (maker,))
        a.rollback()
    with psycopg.connect(URL) as c:
        assert billing.process(c.cursor()) == 2
        c.commit()
    plan, docs, inv = q("""select p.active_plan, (select count(*) from account_documents d where d.user_id = p.user_id),
                           (select status from invoices i where i.user_id = p.user_id) from manufacturer_profiles p where p.user_id = %s""", maker)[0]
    assert plan == "data" and docs == 2 and inv == "paid"           # plan on, order confirmation and paid invoice issued
    with app_conn(env) as a:
        as_user(a, maker)
        assert a.execute("select count(*) from invoices").fetchone()[0] == 1
        pid = a.execute("insert into mfr_products (user_id, name) values (%s, 'X') returning id", (maker,)).fetchone()[0]
        a.execute("insert into launches (user_id, product_id, country, stage) values (%s, %s, 'TZ', 'selected')", (maker, pid))
        assert a.execute("select count(*) from my_market_watch(array['TZ'], 90)").fetchone()[0] >= 0
        a.commit()
        as_user(a, other)
        assert a.execute("select count(*) from launches").fetchone()[0] == 0
        assert a.execute("select count(*) from invoices").fetchone()[0] == 0
        assert a.execute("select count(*) from account_documents").fetchone()[0] == 0
        a.rollback()
        as_user(a, team)                                                     # an invoiced plan, switched on by the team
        a.execute("insert into payment_events (kind, user_id, plan, amount_cents) values ('manual_activation', %s, 'entry', 500000)", (other,))
        a.commit()
    with psycopg.connect(URL) as c:
        billing.process(c.cursor())
        c.execute("""insert into payment_events (stripe_event_id, kind, stripe_subscription) values ('evt_3', 'subscription_ended', 'sub_1')""")
        billing.process(c.cursor())
        c.commit()
    assert q("select status from invoices where user_id = %s", other)[0][0] == "due"
    assert q("select active_plan from manufacturer_profiles where user_id = %s", maker)[0][0] is None   # cancelled subscription ends the plan


def test_team_members_and_permissions(env):
    with psycopg.connect(URL) as c:
        owner = c.execute("insert into users (email, password_hash, role, team_owner) values ('own@example.com', 'x', 'team', true) returning id").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, owner)
        member = a.execute("select team_add_member('mem@example.com', 'Member', 'h', array['requests'], false)").fetchone()[0]
        a.commit()
        with pytest.raises(psycopg.errors.RaiseException):           # unknown permission
            as_user(a, owner)
            a.execute("select team_add_member('x@example.com', 'X', 'h', array['everything'], false)")
        a.rollback()
        as_user(a, member)
        assert a.execute("select app_is_team(), app_is_owner()").fetchone() == (True, False)
        for sql_text, args in (("select team_add_member('y@example.com', 'Y', 'h', array['requests'], true)", ()),   # members cannot add
                               ("select team_set_member(%s, array['accounts'], true)", (member,)),                  # nor promote themselves
                               ("select admin_reset_password(%s, 'h2')", (owner,))):                                # nor take over an owner
            with pytest.raises(psycopg.errors.RaiseException):
                a.execute(sql_text, args)
            a.rollback()
            as_user(a, member)
        a.rollback()
        as_user(a, owner)
        a.execute("select team_set_member(%s, array[]::text[], false)", (owner,))   # owner rights are not changed in the app
        assert q("select team_owner from users where id = %s", owner)[0][0] is True
        with pytest.raises(psycopg.errors.RaiseException):           # and the owner cannot be removed
            a.execute("select team_remove_member(%s, true)", (owner,))
        a.rollback()
        as_user(a, owner)
        a.execute("select team_remove_member(%s, true)", (member,))
        a.commit()
        as_user(a, member)
        assert a.execute("select app_is_team()").fetchone()[0] is False  # removed: no longer team
        assert a.execute("select count(*) from auth_find_user('mem@example.com')").fetchone()[0] == 0   # and cannot sign in
        a.rollback()


def test_files_regulatory_competitor_watch_and_weekly_summary(env):
    import digest
    import notify
    import reg_extract
    tmp, job = env["tmp"], env["job"]
    with psycopg.connect(URL) as c:
        maker = c.execute("insert into users (email, password_hash, account_type, onboarded, whatsapp, notify_whatsapp) values ('cw1@example.com', 'x', 'manufacturer', true, '+255 700 000 001', true) returning id").fetchone()[0]
        other = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('cw2@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        team = c.execute("insert into users (email, password_hash, role) values ('cwt@example.com', 'x', 'team') returning id").fetchone()[0]
        c.execute("""insert into manufacturer_profiles (user_id, company_name, target_countries, active_plan, completed_at)
                     values (%s, 'Watch Pharma', '{TZ}', 'data', now() - interval '8 days'), (%s, 'Other', '{TZ}', null, now())""", (maker, other))
        digoxin = c.execute("select id from molecules where inn = 'Digoxin'").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        # manufacturer files: their own only; the team sees all
        as_user(a, maker)
        fid = a.execute("insert into stored_files (name, size, data, uploaded_by) values ('gmp.pdf', 3, 'pdf', %s) returning id", (maker,)).fetchone()[0]
        a.execute("insert into mfr_files (user_id, kind, file_path, file_name) values (%s, 'gmp', %s, 'gmp.pdf')", (maker, f"db:{fid}"))
        a.execute("insert into watches (user_id, molecule_id) values (%s, %s)", (maker, digoxin))
        a.commit()
        as_user(a, other)
        assert a.execute("select count(*) from mfr_files").fetchone()[0] == 0
        assert a.execute("select count(*) from stored_files where id = %s", (fid,)).fetchone()[0] == 0
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            a.execute("insert into mfr_files (user_id, kind, file_path, file_name) values (%s, 'gmp', 'db:x', 'x')", (maker,))
        a.rollback()
        # regulatory: customers read published requirements only, never the documents; the team writes both
        as_user(a, team)
        did = a.execute("""insert into reg_documents (country, title, file_path, file_name) values ('KZ', 'Guideline', 'db:none', 'g.txt')
                           returning id""").fetchone()[0]
        a.execute("""insert into reg_requirements (country, topic, applies, match, title, body, status, document_id)
                     values ('KZ', 'Dossier', 'molecule', 'digoxin', 'Draft rule', 'Team only for now.', 'draft', %s)""", (did,))
        a.commit()
        as_user(a, other)
        assert a.execute("select count(*) from reg_documents").fetchone()[0] == 0
        titles = [r[0] for r in a.execute("select title from reg_requirements where country = 'KZ'").fetchall()]
        assert "Draft rule" not in titles and "Labels in Kazakh and Russian" in titles
        assert a.execute("select count(*) from reg_requirements where applies = 'form' and match = 'tablet'").fetchone()[0] == 5
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            a.execute("insert into reg_requirements (country, topic, title, body) values ('TZ', 'Dossier', 'x', 'y')")
        a.rollback()
    # a new Digoxin registration: the watching manufacturer gets an alert naming only where it is made
    rows = registry_rows(n=129, extra=[("Doxycycline", "Delta Medics Ltd"), ("Doxycycline", "Beta Distributors (T) Ltd"),
                                       ("Digoxin", "GAMMA HEALTHCARE LTD")], cancel={"TAN 20 HM 0006"})
    job.run(job.create_job(write_registry(tmp / "r3.xls", rows), "upload"))
    alerts = q("select channel, subject, body from notifications where user_id = %s order by channel", maker)
    assert [x[0] for x in alerts] == ["email", "whatsapp"] and alerts[0][1] == "New competitor: Digoxin"
    assert "India" in alerts[0][2] and "GAMMA" not in alerts[0][2].upper() and "MAKER" not in alerts[0][2].upper()
    assert q("select count(*) from notifications where user_id = %s", other)[0][0] == 0
    with app_conn(env) as a:
        as_user(a, maker)
        s = a.execute("select molecule, competitors, new_90, gone_90, last_move is not null, made_in from my_competitor_summary('TZ')").fetchall()
        assert s == [("Digoxin", 1, 1, 0, True, "India")]
        assert a.execute("select change, made_in from my_competitor_watch(array['TZ'], 90)").fetchall() == [("added", "India")]
        a.rollback()
    # weekly summary: once a week, only with something to say
    with psycopg.connect(URL) as c:
        cur = c.cursor()
        c.execute("delete from notifications where user_id = %s", (maker,))
        assert digest.send_weekly(cur) == 1
        assert digest.send_weekly(cur) == 0                     # not again this week
        c.commit()
    body = q("select body from notifications where user_id = %s and channel = 'whatsapp'", maker)[0][0]
    assert "new competitor registration" in body and "GAMMA" not in body.upper()
    assert q("select digest_sent_at is not null from manufacturer_profiles where user_id = %s", other)[0][0] is False  # completed today
    # WhatsApp: template messages carry the text as one variable without line breaks; no credentials means stubbed
    p = notify.whatsapp_payload("+255 700-000 001", "Line one\nLine two", template="nazryx_update")
    assert p["to"] == "255700000001" and p["template"]["components"][0]["parameters"][0]["text"] == "Line one Line two"
    assert notify.send_whatsapp("+255700000001", "hi") == "stubbed"
    # reading documents: only well-formed requirements are kept
    kept = reg_extract.clean([
        {"topic": "Labelling", "applies": "all", "match": "x", "title": "Labels", "body": "Russian labels."},
        {"topic": "Dossier", "applies": "atc", "match": "j01", "title": "Antibiotics", "body": "Extra data."},
        {"topic": "Dossier", "applies": "form", "match": "", "title": "No match", "body": "Dropped."},
        {"topic": "Nonsense", "applies": "all", "match": None, "title": "Bad topic", "body": "Dropped."}])
    assert [(k["applies"], k["match"]) for k in kept] == [("all", None), ("atc", "J01")]


def test_reg_extract_drafts_from_a_document(env):
    import reg_extract

    class Block:
        type = "text"
        text = '{"requirements": [{"topic": "Marking", "applies": "all", "match": null, "title": "Data Matrix", "body": "Every pack."}]}'

    class Client:
        class beta:
            class messages:
                @staticmethod
                def create(**kw):
                    assert kw["messages"][0]["content"][0]["type"] == "text"
                    return type("R", (), {"stop_reason": "end_turn", "content": [Block()]})()

    with psycopg.connect(URL) as c:
        fid = c.execute("insert into stored_files (name, size, data) values ('g.txt', 5, 'rules') returning id").fetchone()[0]
        c.execute("delete from reg_documents")
        did = c.execute("insert into reg_documents (country, title, file_path, file_name) values ('UZ', 'Order', %s, 'g.txt') returning id",
                        (f"db:{fid}",)).fetchone()[0]
        c.commit()
    assert reg_extract.run(client=Client()) == 1
    assert q("select status, found from reg_documents where id = %s", did)[0] == ("done", 1)
    assert q("select status, title from reg_requirements where document_id = %s", did)[0] == ("draft", "Data Matrix")


def test_catalogues_add_their_products(env):
    import io as _io
    import openpyxl
    import catalogue
    with psycopg.connect(URL) as c:
        maker = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('cat1@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        wb = openpyxl.Workbook()
        ws = wb.active
        for r in [["Our products 2026"], ["Product name", "Strength", "Dosage form", "Price"], ["Amoxicillin", "500 mg", "Capsules", "1.20"],
                  ["Atenolol", "50 mg", "Tablets", "0.80"], [None, None, None, None], ["Zzzunknownium", "5 mg", "Tablets", "2"]]:
            ws.append(r)
        buf = _io.BytesIO()
        wb.save(buf)
        ids = []
        for name, data in [("list.xlsx", buf.getvalue()), ("brochure.pdf", b"%PDF-1.4 x")]:
            fid = c.execute("insert into stored_files (name, size, data) values (%s, %s, %s) returning id", (name, len(data), data)).fetchone()[0]
            ids.append(c.execute("""insert into mfr_files (user_id, kind, file_path, file_name, status) values (%s, 'catalogue', %s, %s, 'queued')
                                    returning id""", (maker, f"db:{fid}", name)).fetchone()[0])
        # without an Anthropic credential: Excel is read by its columns, a PDF goes to the team
        assert catalogue.run(c.cursor(), client=None) == 1
        c.commit()
    status = dict(q("select file_name, status from mfr_files where user_id = %s", maker))
    assert status == {"list.xlsx": "done", "brochure.pdf": "manual"}
    assert q("select added, matched from mfr_files where id = %s", ids[0])[0] == (3, 2)
    names = sorted(r[0] for r in q("select name from mfr_products where user_id = %s", maker))
    assert names[0].startswith("Amoxicillin 500 mg") and len(names) == 3            # the title row and the empty row are skipped

    # with Claude: the PDF is read into products; a second read adds no duplicates
    class Block:
        type = "text"
        text = json.dumps({"products": [{"generic": "Ciprofloxacin", "strength": "250 mg", "form": "Tablets", "line": "Ciprofloxacin 250 mg tablets"},
                                        {"generic": "Amoxicillin", "strength": "500 mg", "form": "Capsules", "line": "Amoxicillin 500 mg Capsules"}]})

    class Client:
        class beta:
            class messages:
                @staticmethod
                def create(**kw):
                    assert kw["messages"][0]["content"][0]["type"] == "document"
                    return type("R", (), {"stop_reason": "end_turn", "content": [Block()]})()

    with psycopg.connect(URL) as c:
        c.execute("update mfr_files set status = 'queued' where id = %s", (ids[1],))
        assert catalogue.run(c.cursor(), client=Client()) == 1
        c.commit()
    assert q("select status, added, matched from mfr_files where id = %s", ids[1])[0] == ("done", 1, 1)   # amoxicillin was already there
    assert q("select count(*) from mfr_products where user_id = %s", maker)[0][0] == 4
    assert catalogue.molecule_keyish("Amoxicillin 500mg Capsules BP") == "amoxicillin"

    # a French export: titles, a DCI column next to the brand, several sheets, section headings; read before by the old
    # reader (brand names only, unmatched, and the header of the second sheet saved as a product)
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in [["SOME PHARMA"], ["CATALOGUE DE PRODUITS EXPORT"], [], ["DCI", "TRADE NAME ", "FORM", "PRÉVISIONS VENTES AN 1"],
              ["AMOXICILLINE TRIHYDRATE", "AMOXIL 500 MG B/12", "Capsules"], ["CIPROFLOXACINE", "CIPRO 250 MG B/10", "Tablets"],
              ["OPHTALMOLOGIE  ◆  EYE CARE PRODUCTS"], ["Métformine", "GLUCO 850 MG B/30", "Tablets"], ["DIGOXIN", "DIGO 0,25 MG", "Tablets"]]:
        ws.append(r)
    ws2 = wb.create_sheet("Soft gel")
    for r in [["SOFT GEL  ◆  2026"], ["DCI", "TRADE NAME ", "FORM"], ["PARACETAMOLE", "PARA 500", "Soft Capsule"], ["Vitamin E", "E CAPS 200", "Soft Capsule"]]:
        ws2.append(r)
    buf = _io.BytesIO()
    wb.save(buf)
    with psycopg.connect(URL) as c:
        fr = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('cat2@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        for old in ("AMOXIL 500 MG B/12 Capsules", "TRADE NAME FORM"):
            c.execute("insert into mfr_products (user_id, name) values (%s, %s)", (fr, old))
        sid = c.execute("insert into stored_files (name, size, data) values ('cat.xlsx', %s, %s) returning id", (len(buf.getvalue()), buf.getvalue())).fetchone()[0]
        fid = c.execute("""insert into mfr_files (user_id, kind, file_path, file_name, status) values (%s, 'catalogue', %s, 'cat.xlsx', 'queued')
                           returning id""", (fr, f"db:{sid}")).fetchone()[0]
        assert catalogue.run(c.cursor(), client=None) == 1
        c.commit()
    got = dict(q("""select p.name, split_part(m.inn, ' ', 1) from mfr_products p left join molecules m on m.id = p.molecule_id where p.user_id = %s""", fr))
    assert got == {"AMOXIL 500 MG B/12 Capsules": "Amoxicillin", "CIPRO 250 MG B/10 Tablets": "Ciprofloxacin",
                   "GLUCO 850 MG B/30 Tablets": "Metformin", "DIGO 0,25 MG Tablets": "Digoxin", "PARA 500 Soft Capsule": "Paracetamol",
                   "E CAPS 200 Soft Capsule": None}         # the brand stays the name; 'Vitamin E' is not guessed
    assert q("select added, matched from mfr_files where id = %s", fid)[0] == (6, 5)
    assert q("select dosage_form from mfr_products where user_id = %s and name like 'PARA%%'", fr)[0][0] == "Soft Capsule"   # the form column
    # a refresh fills missing forms on products already there and adds nothing back
    with psycopg.connect(URL) as c:
        c.execute("update mfr_products set dosage_form = null where user_id = %s", (fr,))
        c.execute("delete from mfr_products where user_id = %s and name like 'CIPRO%%'", (fr,))
        c.execute("update mfr_files set status = 'queued', refresh_only = true where id = %s", (fid,))
        assert catalogue.run(c.cursor(), client=None) == 1
        c.commit()
    assert q("select count(*), count(dosage_form) from mfr_products where user_id = %s", fr)[0] == (5, 5)
    assert q("select added, matched, refresh_only from mfr_files where id = %s", fid)[0] == (6, 5, False)


def test_other_markets_load_beside_tanzania(env):
    """Kazakhstan (Russian, ATC codes, no holders) and Kenya (English, local representatives) load through their own
    readers: English names, their own companies, history on the next upload, and Tanzania stays as it was."""
    import openpyxl
    tmp, job = env["tmp"], env["job"]
    tz_before = q("select count(*), count(distinct ltr_id) from registrations where country = 'TZ'")[0]
    tz_companies = q("select count(*) from companies where cluster_key not like '%%|%%'")[0][0]

    def kz_file(name, extra=()):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["№", "Рег. номер", "Торговое название", "Дата регистрации", "Дата истечение", "Производитель", "Страна", "МНН",
                   "Лек. форма", "Формы отпуска", "Код АТХ", "Наименование кода АТХ", "ДРУ", "Адрес ДРУ"])
        mols = [("амоксициллин", "J01CA04", "капсулы 500 мг"), ("азитромицина дигидрат", "J01FA10", "таблетки 500 мг"),
                ("Нет данных", "J01FA10", "порошок для приготовления суспензии"), ("метформина гидрохлорид", "A10BA02", "таблетки 850 мг")]
        for i in range(60):
            inn, atc, form = mols[i % len(mols)]
            ws.append([str(i + 1), f"РК-ЛС-5№{i:06d}", f"Бренд{i}", "01.02.2022", "01.02.2031" if i % 10 else "01.02.2020",
                       f"Завод {i % 5} ООО", "ИНДИЯ" if i % 2 else "КАЗАХСТАН", inn, form, "По рецепту", atc, "", "", ""])
        for j, (inn, atc) in enumerate(extra):
            ws.append([str(100 + j), f"РК-ЛС-5№9{j:05d}", f"Новый{j}", "01.09.2026", "01.09.2031", "Новый завод АО", "ТУРЦИЯ", inn,
                       "таблетки", "По рецепту", atc, "", "", ""])
        wb.save(tmp / name)
        return str(tmp / name)

    first = job.run(job.create_job(kz_file("kz1.xlsx"), "upload", country="KZ"))
    assert first["new_registrations"] == 0                       # the first load of a market is its baseline
    names = dict(q("""select m.inn, cm.registrations from country_molecules cm join molecules m on m.id = cm.molecule_id
                      where cm.country = 'KZ' and cm.registrations > 0"""))
    assert {"Amoxicillin", "Azithromycin"} <= set(names) and any(n.startswith("Metformin") for n in names)   # an existing name is kept
    assert names["Azithromycin"] == 27                            # 'Нет данных' rows named from their ATC code
    assert q("select count(*) from registrations where country = 'KZ' and not active")[0][0] == 6        # expired certificates
    assert q("select count(*) from companies where cluster_key like 'KZ|%%'")[0][0] == 5
    assert q("select distinct manufacturing_country from registrations where country = 'KZ' order by 1") == [("INDIA",), ("KAZAKHSTAN",)]
    second = job.run(job.create_job(kz_file("kz2.xlsx", extra=[("эзомепразол", "A02BC05")]), "upload", country="KZ"))
    assert second["new_registrations"] == 1
    assert q("""select m.inn from registration_history h join registrations r on r.id = h.registration_id
                join molecules m on m.id = r.molecule_id where r.country = 'KZ' and h.change = 'added'""") == [("Esomeprazole",)]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["#", "Product Registration No", "Product Trade Name", "Inn Of Api", "Dosage Form Name", "Country of Origin", "Local Foreign",
               "Mah Company Name", "Local Technical Representative", "↓ Date of Registration", "Date of Expiry"])
    for i in range(60):
        ws.append([str(i + 1), f"H2026/CTD{i:05d}", f"BRAND{i}", ["AMOXICILLIN TRIHYDRATE", "Ciprofloxacin Hydrochloride USP"][i % 2],
                   "Tablet", "India", "Foreign", f"MAKER {i % 4} LTD", ["Alpha Kenya Ltd", "Beta Pharma Kenya"][i % 2], "46294.0", "48120.0"])
    wb.save(tmp / "ke.xlsx")
    job.run(job.create_job(str(tmp / "ke.xlsx"), "upload", country="KE"))
    assert q("select count(distinct ltr_id) from registrations where country = 'KE'")[0][0] == 2
    ke = dict(q("""select m.inn, cm.registrants from country_molecules cm join molecules m on m.id = cm.molecule_id
                   where cm.country = 'KE' and cm.registrations > 0"""))
    assert "Amoxicillin" in ke and "Ciprofloxacin" in ke
    # Kenya's essential list is built in: openings appear even with only two molecules registered
    assert q("select count(*) from country_molecules where country = 'KE' and on_national_list and registrations = 0")[0][0] > 100
    # Tanzania untouched: its products, distributors and companies
    assert q("select count(*), count(distinct ltr_id) from registrations where country = 'TZ'")[0] == tz_before
    assert q("select count(*) from companies where cluster_key not like '%%|%%'")[0][0] == tz_companies
    # one molecule across markets: amoxicillin is the same molecule in Tanzania, Kazakhstan and Kenya
    assert q("""select count(distinct cm.country) from country_molecules cm join molecules m on m.id = cm.molecule_id
                where m.canonical_key = 'amoxicillin' and cm.registrations > 0""")[0][0] == 3
    # partners: Kazakhstan names no distributors, so its list is the companies holding registrations; Kenya lists distributors
    team = q("select id from users where role = 'team' limit 1")[0][0]
    with app_conn(env) as a:
        as_user(a, team)
        kz = a.execute("select by_holder, total, products from partner_search('KZ')").fetchall()
        ke = a.execute("select by_holder, total from partner_search('KE')").fetchall()
        assert kz and all(r[0] for r in kz) and kz[0][1] > 0
        assert ke and not any(r[0] for r in ke) and ke[0][1] == 2
        assert a.execute("select count(*) from partner_search('KZ', null, null, true)").fetchone()[0] < kz[0][1]   # local only
        a.rollback()


def test_pay_per_market_and_yearly_billing(env):
    """A plan covers the market bought with it; an extra market is its own small subscription; data follows the markets."""
    import billing
    with psycopg.connect(URL) as c:
        maker = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('ppm@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        c.execute("insert into manufacturer_profiles (user_id, company_name, target_countries, primary_market) values (%s, 'Markets Pharma', '{TZ,KE}', 'TZ')", (maker,))
        c.commit()
    with app_conn(env) as a:   # the webhook records the yearly plan with Tanzania, then an extra market (Kenya), monthly-priced x10
        a.execute("""insert into payment_events (stripe_event_id, kind, user_id, plan, amount_cents, stripe_customer, stripe_subscription, markets, billing_interval)
                     values ('evt_m1', 'checkout_completed', %s, 'access', 999000, 'cus_m', 'sub_plan', '{TZ}', 'year'),
                            ('evt_m2', 'checkout_completed', %s, 'market', 49000, 'cus_m', 'sub_ke', '{KE}', 'year'),
                            ('evt_m3', 'invoice_paid', %s, 'market', 49000, 'cus_m', 'sub_ke', null, 'year')""", (maker, maker, maker))
        a.commit()
        as_user(a, maker)
        with pytest.raises((psycopg.errors.RaiseException, psycopg.errors.InsufficientPrivilege)):   # a customer cannot add markets to their own plan
            a.execute("update manufacturer_profiles set paid_markets = '{TZ,KE,KZ}' where user_id = %s", (maker,))
        a.rollback()
    with psycopg.connect(URL) as c:
        assert billing.process(c.cursor()) == 3
        c.commit()
    plan, markets, interval = q("select active_plan, paid_markets, billing_interval from manufacturer_profiles where user_id = %s", maker)[0]
    assert plan == "access" and sorted(markets) == ["KE", "TZ"] and interval == "year"
    docs = [r[0] for r in q("select title from account_documents where user_id = %s order by id", maker)]
    assert "Order confirmation: Extra market: Kenya" in docs and any(d.startswith("Invoice NZX-") for d in docs)
    with app_conn(env) as a:
        as_user(a, maker)
        assert a.execute("select mfr_tier_in('TZ'), mfr_tier_in('KE'), mfr_tier_in('KZ')").fetchone() == (2, 2, 0)
        a.rollback()
    with psycopg.connect(URL) as c:   # cancelling Kenya leaves the plan and Tanzania
        c.execute("insert into payment_events (stripe_event_id, kind, stripe_subscription) values ('evt_m4', 'subscription_ended', 'sub_ke')")
        billing.process(c.cursor())
        c.commit()
    assert q("select active_plan, paid_markets from manufacturer_profiles where user_id = %s", maker)[0] == ("access", ["TZ"])
    with app_conn(env) as a:
        as_user(a, maker)
        assert a.execute("select mfr_tier_in('KE')").fetchone()[0] == 0
        a.rollback()


def test_book_a_call(env):
    """One booking per slot, customers see taken times but not who took them, only the team adds links; a reminder once."""
    import digest
    with psycopg.connect(URL) as c:
        a_id = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('call1@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        b_id = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('call2@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        team = c.execute("insert into users (email, password_hash, role) values ('callt@example.com', 'x', 'team') returning id").fetchone()[0]
        slot = c.execute("select date_trunc('hour', now()) + interval '20 hours'").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, a_id)
        a.execute("insert into call_bookings (user_id, starts_at, topic) values (%s, %s, 'plan')", (a_id, slot))
        a.commit()
        as_user(a, b_id)
        assert a.execute("select count(*) from call_bookings").fetchone()[0] == 0                    # not someone else's call
        assert a.execute("select count(*) from taken_call_slots(now(), now() + interval '2 days')").fetchone()[0] == 1
        with pytest.raises(psycopg.errors.UniqueViolation):                                        # one booking per slot
            a.execute("insert into call_bookings (user_id, starts_at, topic) values (%s, %s, 'plan')", (b_id, slot))
        a.rollback()
        as_user(a, a_id)
        with pytest.raises((psycopg.errors.RaiseException, psycopg.errors.InsufficientPrivilege)):  # customers only cancel
            a.execute("update call_bookings set meeting_link = 'https://x.example' where user_id = %s", (a_id,))
        a.rollback()
        as_user(a, team)
        a.execute("update call_bookings set meeting_link = 'https://meet.example/abc' where user_id = %s", (a_id,))
        a.commit()
    with psycopg.connect(URL) as c:
        assert digest.call_reminders(c.cursor()) == 1
        assert digest.call_reminders(c.cursor()) == 0
        c.commit()
    body = q("select body from notifications where user_id = %s and channel = 'whatsapp' order by id desc limit 1", a_id)[0][0]
    assert "https://meet.example/abc" in body


def test_sign_in_and_last_active_are_recorded(env):
    """Signing in records the time; using the portal marks the account active, at most every five minutes."""
    with psycopg.connect(URL) as c:
        uid = c.execute("insert into users (email, password_hash, role) values ('seen@example.com', 'x', 'team') returning id").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        a.execute("select auth_create_session('seen-token', %s, now() + interval '1 day')", (uid,))
        a.commit()
    login, seen = q("select last_login_at, last_seen_at from users where id = %s", uid)[0]
    assert login is not None and seen == login
    with psycopg.connect(URL) as c:
        c.execute("update users set last_seen_at = now() - interval '1 hour' where id = %s", (uid,))
        c.commit()
    with app_conn(env) as a:
        assert a.execute("select auth_session_user('seen-token')").fetchone()[0] == uid
        assert a.execute("select auth_session_user('no-such-token')").fetchone()[0] is None
        a.commit()
    assert q("select last_seen_at > now() - interval '1 minute' from users where id = %s", uid)[0][0] is True


def test_collectors_hold_small_pulls_and_keep_medicine_tenders(env, monkeypatch):
    """A register pull is queued like an upload (the first one as a baseline); a much smaller one waits for the team
    until it is released. Tenders keep medicines only, matched to molecules in English or Russian."""
    import collect
    import storage
    with psycopg.connect(URL) as c:
        cur = c.cursor()
        c.execute("delete from registry_uploads where status = 'queued'")
        monkeypatch.setattr(collect, "REGISTRIES", {"KE": lambda s, tmp: (open(os.path.join(tmp, "k.html"), "w").write("<table></table>") and os.path.join(tmp, "k.html"), 500)})
        monkeypatch.setattr(storage, "store", lambda cur, path, name=None, uploaded_by=None: "db:00000000-0000-0000-0000-000000000000")
        assert collect.run(c, "registry_KE") == "registry_KE"
        assert q("select kind, file_name from registry_uploads where status = 'queued' and country = 'KE'")[0] == ("initial", "collected: k.html")
        c.execute("delete from registry_uploads where status = 'queued'")
        c.execute("update data_sources set last_ok = now() where key = 'registry_KE'")
        c.commit()
        loaded = q("select count(*) from registrations where country = 'KE' and active")[0][0]
        monkeypatch.setattr(collect, "HOLD_RATIO", 10.0 if loaded else 0.0)   # force a hold when anything is loaded
        collect.run(c, "registry_KE")
        held = q("select held_rows, last_error from data_sources where key = 'registry_KE'")[0]
        if loaded:
            assert held[0] == 500 and held[1].startswith("Held:")
            assert not q("select 1 from registry_uploads where status = 'queued' and country = 'KE'")
            c.execute("update data_sources set release_held = true where key = 'registry_KE'")
            c.commit()
            collect.run(c, "registry_KE")
            assert q("select kind from registry_uploads where status = 'queued' and country = 'KE'")[0] == ("upload",)
            assert q("select held_rows, release_held from data_sources where key = 'registry_KE'")[0] == (None, False)
        c.execute("delete from registry_uploads where status = 'queued'")
        c.commit()
        # tenders: a medicine lot (Russian name), a building tender that only names a chemical, an OCDS medicine item
        idx = collect.MoleculeIndex(cur)
        rows = [dict(country="KZ", source="t", ref="1", title="Амоксициллин, таблетка 500 мг", stage="open"),
                dict(country="KZ", source="t", ref="2", title="Supply of building materials", stage="open", text="aluminium"),
                collect._ocds({"ocid": "o-1", "tender": {"title": "Supply of goods", "items": [
                    {"description": "Amoxicillin 500mg caps", "classification": {"id": "51101500"}, "quantity": 1000, "unit": {"value": {"amount": 2.5}}}]},
                    "awards": [{"date": "2026-10-01", "value": {"amount": 2500, "currency": "TZS"}}]}, "TZ", "t", "u"),
                collect._ocds({"ocid": "o-2", "tender": {"title": "Supply of schools laboratory equipment", "items": [
                    {"description": "Glucose 1kg", "classification": {"id": "41000000"}}]}}, "TZ", "t", "u")]
        assert rows[3] is None                                   # not a medicine tender
        assert collect.save_tenders(cur, [r for r in rows if r], idx) == 2
        c.commit()
    got = dict(q("select ref, stage from tenders where source = 't'"))
    assert got == {"1": "open", "o-1": "awarded"}
    amox = q("select id from molecules where canonical_key = 'amoxicillin'")[0][0]
    assert q("select count(*) from tenders where source = 't' and %s = any(molecule_ids)", amox)[0][0] == 2
    assert q("select quantity, unit_price, value from tenders where ref = 'o-1'")[0] == (1000, 2.5, 2500)



def test_tender_text_is_translated_once(env):
    """Russian and Uzbek tender text gets an English line; the same text is asked for once; English text is left alone;
    a changed title drops its old translation."""
    import collect
    import translate
    asked = []

    class Fake:
        class beta:
            class messages:
                @staticmethod
                def create(**kw):
                    texts = [json.loads(ln)["text"] for ln in kw["messages"][0]["content"].splitlines()[1:]]
                    asked.extend(texts)
                    out = {"items": [{"i": i, "en": "EN " + t} for i, t in enumerate(texts)]}
                    return type("R", (), {"stop_reason": "end_turn", "content": [type("B", (), {"type": "text", "text": json.dumps(out)})]})

    with psycopg.connect(URL) as c:
        cur = c.cursor()
        cur.execute("delete from tenders")                    # only this test's tenders (rolled back at the end)
        idx = collect.MoleculeIndex(cur)
        buyer = "Учреждение «Нарынская областная больница»"
        collect.save_tenders(cur, [dict(country="KG", source="tr", ref="1", title="Амоксициллин, таблетки", buyer=buyer, stage="open"),
                                   dict(country="KG", source="tr", ref="2", title="Амоксициллин 500 мг, капсулы", buyer=buyer, stage="open"),
                                   dict(country="UZ", source="tr", ref="3", title="Amoxicillin dori vositalari xaridi", stage="open"),
                                   dict(country="KE", source="tr", ref="4", title="Supply of medicines", stage="open")], idx)
        assert translate.run(cur, None) == 0                   # no key: nothing asked, originals stay
        assert translate.run(cur, Fake) == 4
        assert sorted(asked) == sorted(["Амоксициллин, таблетки", "Амоксициллин 500 мг, капсулы", buyer, "Amoxicillin dori vositalari xaridi"])
        got = dict((r, (t, b)) for r, t, b in cur.execute("select ref, title_en, buyer_en from tenders where source = 'tr'"))
        assert got["1"] == ("EN Амоксициллин, таблетки", "EN " + buyer) and got["2"][1] == "EN " + buyer
        assert got["4"] == (None, None)
        collect.save_tenders(cur, [dict(country="KG", source="tr", ref="1", title="Амоксициллин, капсулы", buyer=buyer, stage="open")], idx)
        assert cur.execute("select title_en, buyer_en from tenders where source = 'tr' and ref = '1'").fetchone() == (None, "EN " + buyer)
        assert translate.run(cur, Fake) == 1 and asked[-1] == "Амоксициллин, капсулы"
        c.rollback()



def test_molecule_prices_from_listings_and_neighbour_estimates(env, monkeypatch):
    """The average per unit uses the best evidence a market has and its usual strength; a market with no price gets
    the neighbour's, adjusted by how the two markets' prices compare on the molecules priced in both."""
    import pricing
    monkeypatch.setattr(pricing, "MIN_OVERLAP", 2)
    monkeypatch.setattr(pricing, "MIN_TYPE", 2)            # price gaps measured from two medicines here (10 in production)
    with psycopg.connect(URL) as c:
        cur = c.cursor()
        for t in ("price_listings", "reference_prices", "tenders"):
            cur.execute(f"delete from {t}")
        mols = [r[0] for r in cur.execute("select id from molecules order by id limit 3")]
        rows = []
        for i, m in enumerate(mols):                  # Kazakhstan is twice Kyrgyzstan's level on the first two
            rows += [("KZ", "a", "cap", f"{m}-1", m, "Таблетки 500 мг №10", 10, 200 * (i + 1) * 4.52, "KZT")] if i < 2 else []
            rows += [("KG", "b", "retail", f"{m}-1", m, "Таблетки 500 мг №10", 10, 100 * (i + 1) * 0.875, "KGS"),
                     ("KG", "b", "retail", f"{m}-2", m, "Таблетки 500 мг №20", 20, 200 * (i + 1) * 0.875, "KGS"),
                     ("KG", "b", "tender", f"{m}-3", m, "Таблетки 500 мг", 1, 1, "KGS"),           # weaker evidence: left out
                     ("KG", "b", "retail", f"{m}-4", m, "Таблетки 250 мг №10", 10, 50, "KGS")]     # another strength
        for r in rows:
            cur.execute("""insert into price_listings (country, source, kind, ref, molecule_id, pack, units, pack_price, currency)
                           values (%s, %s, %s, %s, %s, %s, %s, %s, %s)""", r)
        cur.execute("update fx_rates set per_usd = case currency when 'KZT' then 452 when 'KGS' then 87.5 else per_usd end")
        pricing.rebuild(cur)
        got = {(c_, m): (b, float(u), s, n, f) for c_, m, b, u, s, n, f in
               cur.execute("select country, molecule_id, basis, unit_price, strength, n, factor from molecule_prices")}
        assert got[("KG", mols[0])][:4] == ("listings", 8.75, "500 mg", 2)           # 87.5 and 175 KGS for 10 and 20 tablets
        assert got[("KZ", mols[0])][0] == "listings" and round(got[("KZ", mols[0])][1], 2) == 90.4
        est = got[("KZ", mols[2])]
        assert est[0] == "estimate" and round(float(est[4]), 2) == 2.0 and est[2] == "500 mg"
        assert round(est[1], 1) == round(3 * 8.75 / 87.5 * 2 * 452, 1)              # Kyrgyz price x 2, in tenge
        # a combination nobody prices: its ingredients added up (no measured combination ratio here, so 1.0)
        a, b = [r[0] for r in cur.execute("select inn from molecules where id = any(%s) order by id", (mols[:2],))]
        combo = cur.execute("insert into molecules (inn, canonical_key) values (%s, %s) returning id", (f"{a} + {b}", f"zz combo {a}")).fetchone()[0]
        pricing.rebuild(cur)
        kg = dict(cur.execute("select molecule_id, unit_price from molecule_prices where country = 'KG'").fetchall())
        assert round(float(kg[combo]), 2) == round(float(kg[mols[0]] + kg[mols[1]]), 2)
        c.rollback()



def test_eye_drops_east_africa_from_its_own_shelves_and_south_asia(env):
    """Kenya's one bromfenac bottle on the shelf (no strength on its label) is the market's price, not thrown away; Tanzania,
    with none, takes South Asia and Kenya carried over by the measured eye-drop gaps (about 2x South Asia, 0.29x Kenya)."""
    import pricing
    with psycopg.connect(URL) as c:
        cur = c.cursor()
        for t in ("price_listings", "reference_prices", "tenders"):
            cur.execute(f"delete from {t}")
        m = (cur.execute("select id from molecules where canonical_key = 'bromfenac'").fetchone() or cur.execute(
            "insert into molecules (inn, canonical_key, category) values ('Bromfenac', 'bromfenac', 'medicine') returning id").fetchone())[0]
        rows = [("KE", "Unibrom Eye Drops 5ml", 745, "KES"), ("PK", "Bromfenac (eye) 0.09% Drops 5ml", 380, "PKR")]
        rows += [("IN", f"Brand{i} · packet of 5 ml Eye Drop", p, "INR") for i, p in enumerate((216, 131, 172, 180, 131, 141, 170, 75, 83, 93))]
        rows += [("BD", f"Brand{i} 0.09% Ophthalmic Solution · 5 ml drop", p, "BDT") for i, p in enumerate((100, 100, 100, 110, 110, 100, 115))]
        for i, (co, name, price, ccy) in enumerate(rows):
            cur.execute("""insert into price_listings (country, source, kind, ref, molecule_id, name, pack, units, pack_price, currency)
                           values (%s, 'shop', 'retail', %s, %s, %s, %s, 1, %s, %s)""", (co, f"eye{i}", m, name, name, price, ccy))
        cur.execute("""update fx_rates set per_usd = case currency when 'KES' then 129 when 'TZS' then 2620 when 'PKR' then 281
                       when 'INR' then 88 when 'BDT' then 122 else per_usd end""")
        pricing.rebuild(cur)
        got = {k: float(v) for k, v in cur.execute("select country, unit_price from molecule_prices where molecule_id = %s and unit = 'ml'", (m,))}
        assert round(got["KE"]) == 149                                     # KES 745 for 5 ml, the shelf price
        tz_bottle = got["TZ"] * 5
        assert 4500 < tz_bottle < 7500                                     # not TZS 30,000; near Kenya's 0.29 and South Asia's 2x
        c.rollback()


def test_researched_regulatory_detail_replaces_starter_rules(env, tmp_path):
    """Researched rules replace the generic starter rules (team-written ones stay); low-confidence ones are drafts;
    routes, the document checklist and regulator news pages load; an unchanged file is not reloaded."""
    import regulatory
    (tmp_path / "TZ.json").write_text(json.dumps({"country": "TZ", "checked": "2026-10", "facts": [
        {"topic": "Fees", "applies": "all", "match": None, "title": "Registration fee", "body": "USD 2,000 per product.", "source": "https://x", "confidence": "high"},
        {"topic": "Marking", "applies": "all", "match": None, "title": "No serialization", "body": "None found.", "source": "https://x", "confidence": "low"},
        {"topic": "Dossier", "applies": "form", "match": "sterile", "title": "Sterility", "body": "Validation.", "source": "https://x", "confidence": "high"}],
        "routes": [{"key": "full", "label": "Full", "months_low": 9, "months_high": 18, "fee_usd_low": 2000, "fee_usd_high": 3000},
                   {"key": "abridged", "label": "Abridged", "months_low": 5, "months_high": 10}],
        "documents": [{"name": "Cover letter", "applies": "all"}, {"name": "Sterility validation", "applies": "form", "match": "sterile"}],
        "news": [{"name": "Announcements", "url": "https://www.tmda.go.tz/announcements"}, {"name": "App", "url": "https://x/#/y"}]}))
    with psycopg.connect(URL) as c:
        cur = c.cursor()
        cur.execute("insert into reg_requirements (country, topic, title, body, origin) values ('TZ', 'Other', 'Team note', 'Kept.', 'team')")
        assert regulatory.load(cur, str(tmp_path)) == 1
        got = dict(cur.execute("select title, status from reg_requirements where country = 'TZ' and origin <> 'team'").fetchall())
        assert got == {"Registration fee": "published", "No serialization": "draft", "Sterility": "published"}
        assert cur.execute("select count(*) from reg_requirements where country = 'TZ' and origin = 'team'").fetchone()[0] >= 1
        assert dict(cur.execute("select key, eligibility from country_routes where country = 'TZ'").fetchall()) == {"full": "all", "abridged": "sra_or_pq"}
        assert cur.execute("select count(*) from reg_checklist where country = 'TZ'").fetchone()[0] == 2
        assert [r[0] for r in cur.execute("select url from data_sources where kind = 'news' and country = 'TZ'")] == ["https://www.tmda.go.tz/announcements"]
        assert regulatory.load(cur, str(tmp_path)) == 0                 # unchanged: not reloaded over the team's edits
        c.rollback()


def test_only_the_team_sees_partner_names(env):
    team = q("select id from users where role = 'team' limit 1")[0][0]
    with psycopg.connect(URL) as c:
        cust = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('names@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        c.commit()
    with app_conn(env) as a:
        as_user(a, team)
        names = a.execute("select code, name from team_partner_names('TZ')").fetchall()
        codes = {r[0] for r in a.execute("select code from partner_search('TZ')").fetchall()}
        assert names and codes <= {n[0] for n in names} and all(n[1] for n in names)
        as_user(a, cust)
        assert a.execute("select count(*) from team_partner_names('TZ')").fetchone()[0] == 0
        a.rollback()


def test_manufacturers_can_watch_all_their_molecules(env):
    """The 5-molecule limit was the distributor free plan's; a manufacturer watches every product."""
    with psycopg.connect(URL) as c:
        maker = c.execute("insert into users (email, password_hash, account_type, onboarded) values ('watchall@example.com', 'x', 'manufacturer', true) returning id").fetchone()[0]
        c.commit()
    mols = [r[0] for r in q("select id from molecules order by id limit 8")]
    with app_conn(env) as a:
        as_user(a, maker)
        for m in mols:
            a.execute("insert into watches (user_id, molecule_id) values (%s, %s)", (maker, m))
        a.commit()
    assert q("select count(*) from watches where user_id = %s", maker)[0][0] == len(mols) > 5
